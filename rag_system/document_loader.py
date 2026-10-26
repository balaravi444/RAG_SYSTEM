"""
Document loading, parsing and chunking pipeline.

Supports: PDF, TXT, DOCX, Markdown, CSV, HTML.
"""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import List

from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.documents import Document

from rag_system.config import settings
from rag_system.logger import get_logger

logger = get_logger(__name__)

# ────────────────────────────────────────────────────────────
# Individual file-type loaders
# ────────────────────────────────────────────────────────────

def _load_pdf(path: Path) -> List[Document]:
    """Load a PDF file and return one Document per page."""
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    docs: List[Document] = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            docs.append(
                Document(
                    page_content=text,
                    metadata={"source": str(path), "page": i + 1, "type": "pdf"},
                )
            )
    return docs


def _load_txt(path: Path) -> List[Document]:
    """Load a plain-text file."""
    text = path.read_text(encoding="utf-8", errors="replace")
    return [Document(page_content=text, metadata={"source": str(path), "type": "txt"})]


def _load_docx(path: Path) -> List[Document]:
    """Load a Word .docx file."""
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    full_text = "\n".join(para.text for para in doc.paragraphs if para.text.strip())
    return [Document(page_content=full_text, metadata={"source": str(path), "type": "docx"})]


def _load_markdown(path: Path) -> List[Document]:
    """Load a Markdown file (raw text, preserving structure)."""
    text = path.read_text(encoding="utf-8", errors="replace")
    return [Document(page_content=text, metadata={"source": str(path), "type": "markdown"})]


def _load_csv(path: Path) -> List[Document]:
    """Load a CSV file — each row becomes a Document."""
    text = path.read_text(encoding="utf-8", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    docs: List[Document] = []
    for i, row in enumerate(reader):
        row_text = " | ".join(f"{k}: {v}" for k, v in row.items() if v)
        if row_text.strip():
            docs.append(
                Document(
                    page_content=row_text,
                    metadata={"source": str(path), "row": i + 1, "type": "csv"},
                )
            )
    return docs


def _load_html(path: Path) -> List[Document]:
    """Load an HTML file, stripping tags to plain text."""
    from bs4 import BeautifulSoup

    raw = path.read_text(encoding="utf-8", errors="replace")
    soup = BeautifulSoup(raw, "html.parser")
    text = soup.get_text(separator="\n", strip=True)
    return [Document(page_content=text, metadata={"source": str(path), "type": "html"})]


# ────────────────────────────────────────────────────────────
# Loader registry
# ────────────────────────────────────────────────────────────

_LOADERS = {
    ".pdf": _load_pdf,
    ".txt": _load_txt,
    ".docx": _load_docx,
    ".md": _load_markdown,
    ".csv": _load_csv,
    ".html": _load_html,
    ".htm": _load_html,
}

SUPPORTED_EXTENSIONS = set(_LOADERS.keys())


def load_document(path: Path) -> List[Document]:
    """
    Load a single file into a list of LangChain Documents.

    Raises ``ValueError`` for unsupported file types.
    """
    suffix = path.suffix.lower()
    loader = _LOADERS.get(suffix)
    if loader is None:
        raise ValueError(
            f"Unsupported file type '{suffix}'. "
            f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    logger.info("Loading %s (%s)", path.name, suffix)
    return loader(path)


def load_documents_from_directory(directory: Path) -> List[Document]:
    """
    Recursively load all supported files from a directory.
    """
    all_docs: List[Document] = []
    files = sorted(
        f for f in directory.rglob("*") if f.is_file() and f.suffix.lower() in SUPPORTED_EXTENSIONS
    )

    if not files:
        logger.warning("No supported files found in %s", directory)
        return all_docs

    logger.info("Found %d supported files in %s", len(files), directory)

    for filepath in files:
        try:
            docs = load_document(filepath)
            all_docs.extend(docs)
        except Exception as exc:
            logger.error("Failed to load %s: %s", filepath, exc)

    logger.info("Loaded %d raw document sections total", len(all_docs))
    return all_docs


# ────────────────────────────────────────────────────────────
# Chunking
# ────────────────────────────────────────────────────────────

def chunk_documents(documents: List[Document]) -> List[Document]:
    """
    Split documents into smaller overlapping chunks for embedding.
    """
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
        length_function=len,
        separators=["\n\n", "\n", ". ", " ", ""],
    )

    chunks = splitter.split_documents(documents)
    logger.info(
        "Chunked %d documents → %d chunks (size=%d, overlap=%d)",
        len(documents),
        len(chunks),
        settings.chunk_size,
        settings.chunk_overlap,
    )
    return chunks
